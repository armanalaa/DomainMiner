"""
extract_erds.py
===============
Generate ERD diagrams for datasets with completed tuning results.

Outputs:
  ERDs/<DatasetName>/full_schema.png
  ERDs/<DatasetName>/domains/D<id>_<DomainName>.png
  ERDs/<DatasetName>/inferred_relationships.txt

The upstream schema.json stores only "PK", "FK", or "" in each column's key
field. It does not store exact FK targets, so relationship lines are inferred
from column names, table/entity names, and generated descriptions.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
except ImportError:
    sys.exit("Missing: pip install openpyxl")

try:
    from sqlalchemy import (
        Boolean,
        Column,
        DateTime,
        Float,
        ForeignKey,
        Integer,
        MetaData,
        String,
        Table,
        Text,
        create_engine,
    )
    from eralchemy2 import render_er
except ImportError:
    sys.exit("Missing: pip install eralchemy2 sqlalchemy")


UNSAFE_NAME_RE = re.compile(r"[^A-Za-z0-9_]")
FK_SUFFIXES = ("_id", "id", "_fk", "fk", "_key", "key", "_code", "code")


def check_graphviz() -> None:
    """Fail fast if Graphviz dot is unavailable."""
    import shutil
    import subprocess

    dot_path = shutil.which("dot")
    if not dot_path:
        sys.exit(
            "\n[FATAL] Graphviz 'dot' executable not found on PATH.\n"
            "Install Graphviz, add its bin folder to PATH, then open a new terminal.\n"
            "Verify with: dot -V\n"
        )
    try:
        subprocess.run([dot_path, "-V"], capture_output=True, timeout=10, check=True)
    except Exception as exc:
        sys.exit(f"\n[FATAL] Found dot at {dot_path}, but it failed to run: {exc}\n")


def _sa_type(sql_type: str):
    """Map a SQL type string from schema.json to a SQLAlchemy type instance."""
    t = (sql_type or "").upper().strip()
    if any(t.startswith(p) for p in ("INT", "TINYINT", "SMALLINT", "MEDIUMINT", "BIGINT", "SERIAL", "ROWID")):
        return Integer()
    if any(t.startswith(p) for p in ("FLOAT", "DOUBLE", "REAL", "NUMERIC", "DECIMAL", "NUMBER")):
        return Float()
    if t.startswith("BOOL"):
        return Boolean()
    if any(t.startswith(p) for p in ("DATE", "TIME", "TIMESTAMP", "DATETIME", "YEAR")):
        return DateTime()
    if any(t.startswith(p) for p in ("TEXT", "CLOB", "BLOB", "MEDIUMTEXT", "LONGTEXT")):
        return Text()
    return String()


def _slug(name: str) -> str:
    return re.sub(r"[^\w\-]", "_", name).strip("_")[:60] or "unnamed"


def _safe_identifier(raw: str, fallback: str) -> str:
    raw = (raw or "").strip()
    if not raw:
        return fallback
    safe = UNSAFE_NAME_RE.sub("_", raw)
    if safe and safe[0].isdigit():
        safe = f"_{safe}"
    return safe or fallback


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def _split_words(name: str) -> list[str]:
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", name or "")
    return re.findall(r"[a-z0-9]+", text.lower())


def _singular(word: str) -> str:
    if len(word) > 3 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith(("ses", "xes", "ches", "shes")):
        return word[:-2]
    if len(word) > 2 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


def _entity_norm(name: str, strip_namespace: bool = True) -> str:
    words = [_singular(w) for w in _split_words(name)]
    if strip_namespace and len(words) > 2 and words[:2] == ["human", "resource"]:
        words = words[2:]
    elif strip_namespace and len(words) > 1 and words[0] in {
        "person",
        "production",
        "purchasing",
        "sale",
        "sales",
        "dbo",
        "public",
    }:
        words = words[1:]
    words = [w for w in words if w not in {"history", "archive", "detail", "header"}]
    return "".join(words)


def _column_entity_norm(column_name: str) -> str:
    words = [_singular(w) for w in _split_words(column_name)]
    while words and words[0] in {"bill", "ship", "to", "from"}:
        words.pop(0)
    if words and words[-1] == "id":
        words.pop()
    return "".join(words)


def _entity_hint(description: str) -> str:
    match = re.search(
        r"(?:foreign key )?(?:linking|links|references?)\s+(?:to\s+)?"
        r"(?:the\s+)?(.+?)(?:\s+entity|\s+table|\.|$)",
        description or "",
        flags=re.IGNORECASE,
    )
    return match.group(1).strip() if match else ""


def _raw_tables(schema: dict) -> dict:
    tables = schema.get("tables", schema)
    if isinstance(tables, list):
        return {t.get("name", str(i)): t for i, t in enumerate(tables)}
    return tables


def build_table_lookup(raw_tables: dict) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for tname in raw_tables:
        sing = _singular(tname)
        entity = _entity_norm(tname)
        for variant in (tname, sing, f"{tname}s", f"{sing}s", entity, f"{entity}s"):
            norm = _normalize(variant)
            if norm and norm not in lookup:
                lookup[norm] = tname
    return lookup


def build_pk_map(raw_tables: dict) -> dict[str, str]:
    """Map table name to its primary key, falling back to first FK-like key."""
    pk_map: dict[str, str] = {}
    for tname, tdata in raw_tables.items():
        for col in tdata.get("columns", []):
            if col.get("key") == "PK":
                cname = (col.get("name") or "").strip()
                if cname:
                    pk_map[tname] = cname
                    break
        if tname in pk_map:
            continue
        for col in tdata.get("columns", []):
            if col.get("key") == "FK":
                cname = (col.get("name") or "").strip()
                if cname:
                    pk_map[tname] = cname
                    break
    return pk_map


def infer_fk_target(
    col_name: str,
    description: str,
    table_lookup: dict[str, str],
    pk_map: dict[str, str],
) -> tuple[str, str | None, str | None]:
    """Return (status, target_table, target_col) for a probable FK target."""
    candidates: list[str] = []
    hint = _entity_norm(_entity_hint(description), strip_namespace=False)
    col_entity = _column_entity_norm(col_name)
    if hint:
        candidates.append(hint)
    if col_entity:
        candidates.append(col_entity)

    base = (col_name or "").strip()
    base_lower = base.lower()
    for suffix in sorted(FK_SUFFIXES, key=len, reverse=True):
        if base_lower.endswith(suffix) and len(base_lower) > len(suffix):
            stripped = base[: -len(suffix)].rstrip("_")
            if stripped:
                candidates.append(stripped)
    candidates.append(base)

    seen: set[str] = set()
    for candidate in candidates:
        norm = _normalize(candidate)
        if not norm or norm in seen:
            continue
        seen.add(norm)
        target_table = table_lookup.get(norm)
        if not target_table:
            continue
        target_col = pk_map.get(target_table)
        if target_col:
            return "resolved", target_table, target_col
        return "no_pk", target_table, None
    return "unmatched", None, None


def build_metadata(
    schema: dict,
    table_filter: list[str] | None = None,
    infer_fks: bool = True,
) -> tuple[MetaData, list[tuple[str, str, str, str]], list[tuple[str, str]], list[tuple[str, str, str]]]:
    """Build SQLAlchemy metadata from schema.json."""
    engine = create_engine("sqlite:///:memory:")
    meta = MetaData()
    raw_tables = _raw_tables(schema)
    included_tables = set(table_filter) if table_filter else set(raw_tables.keys())

    table_lookup = build_table_lookup(raw_tables) if infer_fks else {}
    pk_map = build_pk_map(raw_tables) if infer_fks else {}

    safe_table_name: dict[str, str] = {}
    for idx, tname in enumerate(raw_tables.keys()):
        safe = _safe_identifier(tname, f"_table_{idx}")
        if safe != tname:
            print(f"    [WARN] table name '{tname}' sanitized as '{safe}'")
        safe_table_name[tname] = safe

    inferred_fks: list[tuple[str, str, str, str]] = []
    unresolved_fks: list[tuple[str, str]] = []
    no_pk_fks: list[tuple[str, str, str]] = []

    for tname, tdata in raw_tables.items():
        if tname not in included_tables:
            continue

        cols: list[Column] = []
        seen_names: set[str] = set()
        pk_count = sum(1 for c in tdata.get("columns", []) if c.get("key", "") == "PK")

        for col_idx, col in enumerate(tdata.get("columns", [])):
            raw_cname = (col.get("name") or "").strip()
            if not raw_cname:
                raw_cname = f"_unnamed_col_{col_idx}"
                print(f"    [WARN] {tname}: blank column #{col_idx}; using '{raw_cname}'")

            cname = _safe_identifier(raw_cname, f"_col_{col_idx}")
            base_cname = cname
            dedup_i = 1
            while cname in seen_names:
                dedup_i += 1
                cname = f"{base_cname}_{dedup_i}"
            seen_names.add(cname)

            sql_t = col.get("sql_type", "TEXT")
            key = col.get("key", "")
            sa_type = _sa_type(sql_t)
            description = col.get("description", "")

            if key == "PK":
                source_entity = _entity_norm(tname)
                column_entity = _column_entity_norm(raw_cname)
                pk_is_relationship = (
                    infer_fks
                    and pk_count > 1
                    and column_entity
                    and column_entity != source_entity
                    and source_entity not in column_entity
                )
                if pk_is_relationship:
                    status, target_table, target_col = infer_fk_target(raw_cname, description, table_lookup, pk_map)
                    if status == "resolved" and target_table in included_tables:
                        safe_ttable = safe_table_name[target_table]
                        safe_tcol = _safe_identifier(target_col, target_col)
                        try:
                            cols.append(Column(cname, sa_type, ForeignKey(f"{safe_ttable}.{safe_tcol}"), primary_key=True))
                            inferred_fks.append((tname, cname, target_table, target_col))
                            continue
                        except Exception:
                            pass
                cols.append(Column(cname, sa_type, primary_key=True))
                continue

            if key == "FK" and infer_fks:
                status, target_table, target_col = infer_fk_target(raw_cname, description, table_lookup, pk_map)
                if status == "resolved" and target_table in included_tables:
                    safe_ttable = safe_table_name[target_table]
                    safe_tcol = _safe_identifier(target_col, target_col)
                    try:
                        cols.append(Column(cname, sa_type, ForeignKey(f"{safe_ttable}.{safe_tcol}")))
                        inferred_fks.append((tname, cname, target_table, target_col))
                        continue
                    except Exception:
                        pass
                elif status == "no_pk":
                    no_pk_fks.append((tname, cname, target_table))
                elif status == "unmatched":
                    unresolved_fks.append((tname, cname))

            cols.append(Column(cname, sa_type))

        if cols:
            try:
                Table(safe_table_name[tname], meta, *cols)
            except Exception as exc:
                print(f"    [WARN] {tname}: skipped; could not build table: {exc}")

    try:
        meta.create_all(engine)
    except Exception:
        pass

    return meta, inferred_fks, unresolved_fks, no_pk_fks


def read_best_run_tag(dataset_dir: Path) -> str | None:
    xlsx = dataset_dir / "ccm_output" / "tune_params_results.xlsx"
    if not xlsx.exists():
        return None
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        ws = wb.active
        row = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
        return str(row[3]) if row[3] is not None else None
    except Exception as exc:
        print(f"  [WARN] Could not read {xlsx}: {exc}")
        return None


def render_erd(meta: MetaData, out_path: Path, title: str, fmt: str) -> bool:
    if not list(meta.tables):
        print(f"    [SKIP] No tables; skipping {out_path.name}")
        return False

    target = str(out_path.with_suffix(f".{fmt}"))
    try:
        render_er(meta, target, title=title)
        if not Path(target).exists():
            print(f"    [ERR] {out_path.name}: render_er returned, but no file was written")
            return False
        print(f"    OK {out_path.name} ({len(list(meta.tables))} tables)")
        return True
    except Exception as exc:
        import traceback

        print(f"    [ERR] {out_path.name}: {type(exc).__name__}: {exc}")
        traceback.print_exc()
        return False


def process_dataset(dataset_dir: Path, erd_root: Path, fmt: str, infer_fks: bool = True) -> tuple[int, int]:
    name = dataset_dir.name
    print(f"\n  {'-' * 50}")
    print(f"  Dataset: {name}")

    schema_path = dataset_dir / "schema.json"
    if not schema_path.exists():
        print("  [SKIP] schema.json not found")
        return (0, 0)
    try:
        schema = json.loads(schema_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception as exc:
        print(f"  [SKIP] Cannot read schema.json: {exc}")
        return (0, 0)

    run_tag = read_best_run_tag(dataset_dir)
    if not run_tag:
        print("  [SKIP] tune_params_results.xlsx missing or empty")
        return (0, 0)

    domains_path = dataset_dir / "ccm_output" / run_tag / "step5_domains.json"
    if not domains_path.exists():
        print(f"  [SKIP] step5_domains.json not found at {domains_path}")
        return (0, 0)
    try:
        domains = json.loads(domains_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception as exc:
        print(f"  [SKIP] Cannot read step5_domains.json: {exc}")
        return (0, 0)

    print(f"  Best run : {run_tag}")
    print(f"  Domains  : {len(domains)}")

    dataset_out = erd_root / name
    domains_out = dataset_out / "domains"
    dataset_out.mkdir(parents=True, exist_ok=True)
    domains_out.mkdir(parents=True, exist_ok=True)

    n_ok, n_fail = 0, 0
    all_inferred: list[tuple[str, str, str, str]] = []
    all_unresolved: list[tuple[str, str]] = []
    all_no_pk: list[tuple[str, str, str]] = []

    print("  Rendering full-schema ERD ...")
    full_meta, inf, unr, no_pk = build_metadata(schema, infer_fks=infer_fks)
    all_inferred.extend(inf)
    all_unresolved.extend(unr)
    all_no_pk.extend(no_pk)
    if infer_fks:
        print(
            f"    FK relationships inferred: {len(inf)} "
            f"(unresolved: {len(unr)}, no-PK-on-target: {len(no_pk)})"
        )
    if render_erd(full_meta, dataset_out / "full_schema", f"{name} - Full Schema ({len(full_meta.tables)} tables)", fmt):
        n_ok += 1
    else:
        n_fail += 1

    print(f"  Rendering {len(domains)} domain ERD(s) ...")
    for domain in domains:
        did = domain.get("domain_id", "?")
        dname = domain.get("domain_name", f"Domain_{did}")
        tables = domain.get("tables", [])
        if not tables:
            continue
        domain_meta, _, _, _ = build_metadata(schema, table_filter=tables, infer_fks=infer_fks)
        fname = f"D{did}_{_slug(dname)}"
        title = f"{name} - D{did}: {dname} ({len(tables)} tables)"
        if render_erd(domain_meta, domains_out / fname, title, fmt):
            n_ok += 1
        else:
            n_fail += 1

    if infer_fks:
        log_path = dataset_out / "inferred_relationships.txt"
        lines = [
            f"FK relationship inference for {name}",
            "(heuristic: inferred from schema.json names/descriptions; exact FK targets were not stored)",
            "=" * 70,
            "",
            f"Inferred ({len(all_inferred)}):",
        ]
        for tname, cname, ttable, tcol in sorted(set(all_inferred)):
            lines.append(f"  {tname}.{cname} -> {ttable}.{tcol}")
        lines += ["", f"Matched target table but no target PK-like column ({len(all_no_pk)}):"]
        for tname, cname, ttable in sorted(set(all_no_pk)):
            lines.append(f"  {tname}.{cname} -> {ttable}")
        lines += ["", f"Unresolved ({len(all_unresolved)}):"]
        for tname, cname in sorted(set(all_unresolved)):
            lines.append(f"  {tname}.{cname}")
        log_path.write_text("\n".join(lines), encoding="utf-8")

    if n_fail:
        print(f"  WARN {name}: {n_ok} rendered, {n_fail} failed")
    return (n_ok, n_fail)


def discover_datasets(root: Path) -> list[str]:
    found = []
    for child in sorted(root.iterdir()):
        if child.is_dir() and (child / "ccm_output" / "tune_params_results.xlsx").exists():
            found.append(child.name)
    return found


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate full-schema and per-domain ERD diagrams.")
    parser.add_argument("--root", type=Path, default=Path("."), help="Project root (default: current directory)")
    parser.add_argument("--output_dir", type=Path, default=None, help="ERD output root (default: <root>/ERDs)")
    parser.add_argument("--datasets", nargs="+", default=None, help="Specific datasets to process")
    parser.add_argument("--format", choices=["png", "pdf", "svg"], default="png", help="Output format")
    parser.add_argument("--no-infer-fks", action="store_true", help="Do not infer FK relationship lines")
    args = parser.parse_args()

    root = args.root.resolve()
    erd_root = (args.output_dir or root / "ERDs").resolve()
    datasets = args.datasets or discover_datasets(root)
    infer_fks = not args.no_infer_fks

    print("\nDomainMiner - ERD Generator")
    print(f"Root       : {root}")
    print(f"ERD output : {erd_root}")
    print(f"Format     : {args.format}")
    print(f"FK lines   : {'inferred' if infer_fks else 'disabled'}")
    print(f"Datasets   : {len(datasets)} candidate(s)")

    check_graphviz()
    erd_root.mkdir(parents=True, exist_ok=True)

    total_ok, total_fail, skipped = 0, 0, 0
    for name in datasets:
        dataset_dir = root / name
        if not dataset_dir.is_dir():
            print(f"\n  [SKIP] {name}: folder not found")
            skipped += 1
            continue
        try:
            n_ok, n_fail = process_dataset(dataset_dir, erd_root, args.format, infer_fks=infer_fks)
            total_ok += n_ok
            total_fail += n_fail
            if n_ok == 0 and n_fail == 0:
                skipped += 1
        except Exception as exc:
            import traceback

            print(f"\n  [ERR] {name}: {exc}")
            traceback.print_exc()
            skipped += 1

    print(f"\n{'-' * 52}")
    print(f"Done. {total_ok} ERD(s) rendered, {total_fail} failed, {skipped} dataset(s) skipped.")
    print(f"ERDs written to: {erd_root}")


if __name__ == "__main__":
    main()
