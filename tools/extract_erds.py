"""
extract_erds.py
===============
Generate full-schema and per-domain ERDs directly from dataset CSV files.

The ERD table/column structure is read from:
  <Dataset>/csv/*.csv

Relationship lines are inferred from the CSV values themselves:
  if the distinct non-null values in a source column are contained in a unique
  target column, and the names are compatible, an FK edge is drawn.

Domain membership is still read from the pipeline output:
  <Dataset>/ccm_output/<best_run>/step5_domains.json
because that is where the discovered table groups are stored.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PIPELINE_DIR = PROJECT_ROOT / "pipeline"
if str(PIPELINE_DIR) not in sys.path:
    sys.path.insert(0, str(PIPELINE_DIR))

from path_utils import DATALAKES_ROOT

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
except ImportError:
    sys.exit("Missing: pip install openpyxl")

try:
    from sqlalchemy import (
        Boolean,
        Column,
        DateTime,
        Float,
        ForeignKey,
        Integer,
        MetaData,
        String,
        Table,
        Text,
        create_engine,
    )
    from eralchemy2 import render_er
except ImportError:
    sys.exit("Missing: pip install eralchemy2 sqlalchemy")


UNSAFE_NAME_RE = re.compile(r"[^A-Za-z0-9_]")
KEYLIKE_RE = re.compile(r"(id|code)$", re.IGNORECASE)
DEFAULT_MAX_DISTINCT = 200_000


@dataclass
class CsvColumn:
    name: str
    safe_name: str
    sql_type: str
    values: set[str] = field(default_factory=set)
    non_null_count: int = 0
    capped: bool = False

    @property
    def is_unique(self) -> bool:
        return self.non_null_count > 0 and not self.capped and len(self.values) == self.non_null_count


@dataclass
class CsvTable:
    name: str
    safe_name: str
    columns: list[CsvColumn]
    row_count: int


def check_graphviz() -> None:
    import shutil
    import subprocess

    dot_path = shutil.which("dot")
    if not dot_path:
        sys.exit(
            "\n[FATAL] Graphviz 'dot' executable not found on PATH.\n"
            "Install Graphviz, add its bin folder to PATH, then open a new terminal.\n"
            "Verify with: dot -V\n"
        )
    try:
        subprocess.run([dot_path, "-V"], capture_output=True, timeout=10, check=True)
    except Exception as exc:
        sys.exit(f"\n[FATAL] Found dot at {dot_path}, but it failed to run: {exc}\n")


def _slug(name: str) -> str:
    return re.sub(r"[^\w\-]", "_", name).strip("_")[:60] or "unnamed"


def _safe_identifier(raw: str, fallback: str) -> str:
    raw = (raw or "").strip()
    if not raw:
        return fallback
    safe = UNSAFE_NAME_RE.sub("_", raw)
    if safe and safe[0].isdigit():
        safe = f"_{safe}"
    return safe or fallback


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def _split_words(name: str) -> list[str]:
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", name or "")
    return re.findall(r"[a-z0-9]+", text.lower())


def _singular(word: str) -> str:
    if len(word) > 3 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith(("ses", "xes", "ches", "shes")):
        return word[:-2]
    if len(word) > 2 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


def _entity_norm(name: str, strip_namespace: bool = True) -> str:
    words = [_singular(w) for w in _split_words(name)]
    if strip_namespace and len(words) > 2 and words[:2] == ["human", "resource"]:
        words = words[2:]
    elif strip_namespace and len(words) > 1 and words[0] in {
        "person",
        "production",
        "purchasing",
        "sale",
        "sales",
        "dbo",
        "public",
    }:
        words = words[1:]
    words = [w for w in words if w not in {"history", "archive"}]
    return "".join(words)


def _canonical_entity(value: str) -> str:
    aliases = {
        "prod": "product",
        "product": "product",
        "cust": "customer",
        "customer": "customer",
        "inv": "inventory",
        "inventory": "inventory",
    }
    return aliases.get(value, value)


def _column_entity_norm(column_name: str) -> str:
    words = [_singular(w) for w in _split_words(column_name)]
    while words and words[0] in {"bill", "ship", "to", "from"}:
        words.pop(0)
    if words and words[-1] in {"id", "code"}:
        words.pop()
    entity = "".join(words)
    if entity.endswith("id") and len(entity) > 2:
        entity = entity[:-2]
    elif entity.endswith("code") and len(entity) > 4:
        entity = entity[:-4]
    return _canonical_entity(entity)


def _is_null(value: str | None) -> bool:
    if value is None:
        return True
    return str(value).strip().upper() in {"", "NULL", "NONE", "NAN"}


def _is_int(value: str) -> bool:
    return bool(re.fullmatch(r"[+-]?\d+", value.strip()))


def _is_float(value: str) -> bool:
    return bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value.strip()))


def _is_bool(value: str) -> bool:
    return value.strip().lower() in {"true", "false", "0", "1", "yes", "no"}


def _looks_datetime(value: str) -> bool:
    return bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}:\d{2}(?:\.\d+)?)?", value.strip()))


def infer_sql_type(values: list[str]) -> str:
    clean = [v for v in values if not _is_null(v)]
    if not clean:
        return "VARCHAR(50)"
    if all(_is_bool(v) for v in clean):
        return "BOOLEAN"
    if all(_is_int(v) for v in clean):
        return "INT"
    if all(_is_float(v) for v in clean):
        return "FLOAT"
    if all(_looks_datetime(v) for v in clean):
        return "DATETIME"
    max_len = max(len(str(v)) for v in clean)
    if max_len <= 50:
        return "VARCHAR(50)"
    if max_len <= 255:
        return "VARCHAR(255)"
    return "TEXT"


def _sa_type(sql_type: str):
    t = (sql_type or "").upper().strip()
    if t.startswith("INT"):
        return Integer()
    if t.startswith("FLOAT"):
        return Float()
    if t.startswith("BOOL"):
        return Boolean()
    if t.startswith(("DATE", "TIME")):
        return DateTime()
    if t.startswith("TEXT"):
        return Text()
    return String()


def is_keylike_column(name: str) -> bool:
    lower = name.lower()
    return lower != "rowguid" and bool(KEYLIKE_RE.search(lower))


def read_csv_tables(csv_dir: Path, max_distinct: int = DEFAULT_MAX_DISTINCT) -> dict[str, CsvTable]:
    if not csv_dir.exists():
        raise FileNotFoundError(f"CSV folder not found: {csv_dir}")

    tables: dict[str, CsvTable] = {}
    for table_idx, path in enumerate(sorted(csv_dir.glob("*.csv"))):
        tname = path.stem
        safe_tname = _safe_identifier(tname, f"_table_{table_idx}")
        if safe_tname != tname:
            print(f"    [WARN] table name '{tname}' sanitized as '{safe_tname}'")

        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            try:
                header = next(reader)
            except StopIteration:
                print(f"    [WARN] {path.name}: empty CSV; skipped")
                continue

            seen: set[str] = set()
            raw_to_safe: list[tuple[str, str]] = []
            for col_idx, raw_col in enumerate(header):
                raw_col = (raw_col or "").strip() or f"_unnamed_col_{col_idx}"
                safe_col = _safe_identifier(raw_col, f"_col_{col_idx}")
                base = safe_col
                dedup_i = 1
                while safe_col in seen:
                    dedup_i += 1
                    safe_col = f"{base}_{dedup_i}"
                seen.add(safe_col)
                raw_to_safe.append((raw_col, safe_col))

            samples: dict[str, list[str]] = {raw: [] for raw, _ in raw_to_safe}
            values: dict[str, set[str]] = {raw: set() for raw, _ in raw_to_safe}
            non_null_count: dict[str, int] = {raw: 0 for raw, _ in raw_to_safe}
            capped: dict[str, bool] = {raw: False for raw, _ in raw_to_safe}
            row_count = 0

            for row in reader:
                row_count += 1
                if len(row) < len(raw_to_safe):
                    row = row + [""] * (len(raw_to_safe) - len(row))
                for idx, (raw_col, _) in enumerate(raw_to_safe):
                    value = row[idx] if idx < len(row) else ""
                    if len(samples[raw_col]) < 100 and not _is_null(value):
                        samples[raw_col].append(value)
                    if _is_null(value):
                        continue
                    non_null_count[raw_col] += 1
                    if not capped[raw_col]:
                        values[raw_col].add(str(value))
                        if len(values[raw_col]) > max_distinct:
                            values[raw_col].clear()
                            capped[raw_col] = True

        columns = [
            CsvColumn(
                name=raw_col,
                safe_name=safe_col,
                sql_type=infer_sql_type(samples[raw_col]),
                values=values[raw_col],
                non_null_count=non_null_count[raw_col],
                capped=capped[raw_col],
            )
            for raw_col, safe_col in raw_to_safe
        ]
        tables[tname] = CsvTable(tname, safe_tname, columns, row_count)

    return tables


def choose_primary_columns(tables: dict[str, CsvTable]) -> dict[str, str]:
    """Choose one display primary key per table from unique CSV columns."""
    pk_map: dict[str, str] = {}
    for tname, table in tables.items():
        entity = _entity_norm(tname)
        unique_keylike = [c for c in table.columns if c.is_unique and is_keylike_column(c.name)]
        if not unique_keylike:
            continue

        def score(col: CsvColumn) -> tuple[int, int]:
            col_norm = _normalize(col.name)
            col_entity = _column_entity_norm(col.name)
            s = 0
            if col_entity and col_entity == entity:
                s += 100
            if entity and entity in col_norm:
                s += 60
            if col.name.lower().endswith("id"):
                s += 30
            if col.name.lower().endswith("code"):
                s += 20
            return s, -len(col.name)

        best = max(unique_keylike, key=score)
        pk_map[tname] = best.name
    return pk_map


def relationship_score(
    source_table: str,
    source_col: CsvColumn,
    target_table: str,
    target_col: CsvColumn,
) -> int:
    if source_table == target_table and source_col.name == target_col.name:
        return 0
    if not source_col.values or source_col.capped:
        return 0
    if not target_col.is_unique or not target_col.values:
        return 0
    if not source_col.values.issubset(target_col.values):
        return 0

    source_name = source_col.name
    target_name = target_col.name
    source_entity = _column_entity_norm(source_name)
    target_entity = _entity_norm(target_table)
    source_table_entity = _entity_norm(source_table)
    source_norm = _normalize(source_name)
    target_col_norm = _normalize(target_name)
    target_table_norm = _normalize(target_table)

    if source_entity == source_table_entity and source_entity != target_entity:
        return 0
    if source_col.is_unique and source_entity != target_entity:
        return 0

    score = 0
    if source_table == target_table and source_col.name != target_col.name:
        score += 15
    if source_name == target_name:
        score += 60
    elif target_col_norm and target_col_norm in source_norm:
        score += 45
    if source_entity and source_entity == target_entity:
        score += 80
    elif source_entity and (source_entity in target_table_norm or target_entity in source_entity):
        score += 35
    if is_keylike_column(source_name):
        score += 20
    if is_keylike_column(target_name):
        score += 20
    if source_col.sql_type == target_col.sql_type:
        score += 10
    return score


def infer_relationships(
    tables: dict[str, CsvTable],
    included_tables: set[str],
    min_score: int = 70,
) -> tuple[list[tuple[str, str, str, str]], list[tuple[str, str]]]:
    """Infer FK edges from CSV value containment among included tables."""
    target_cols: list[tuple[str, CsvColumn]] = []
    for tname in included_tables:
        table = tables.get(tname)
        if not table:
            continue
        for col in table.columns:
            if col.is_unique and is_keylike_column(col.name):
                target_cols.append((tname, col))

    inferred: list[tuple[str, str, str, str]] = []
    unresolved: list[tuple[str, str]] = []
    seen_edges: set[tuple[str, str, str, str]] = set()

    for source_table in sorted(included_tables):
        table = tables.get(source_table)
        if not table:
            continue
        for source_col in table.columns:
            if not is_keylike_column(source_col.name):
                continue

            candidates: list[tuple[int, str, CsvColumn]] = []
            for target_table, target_col in target_cols:
                score = relationship_score(source_table, source_col, target_table, target_col)
                if score >= min_score:
                    candidates.append((score, target_table, target_col))

            if not candidates:
                unresolved.append((source_table, source_col.name))
                continue

            candidates.sort(key=lambda x: (x[0], x[2].non_null_count), reverse=True)
            best_score, target_table, target_col = candidates[0]
            if len(candidates) > 1 and candidates[1][0] == best_score:
                unresolved.append((source_table, source_col.name))
                continue
            edge = (source_table, source_col.name, target_table, target_col.name)
            if edge not in seen_edges:
                inferred.append(edge)
                seen_edges.add(edge)

    return inferred, unresolved


def build_metadata_from_csv(
    tables: dict[str, CsvTable],
    table_filter: list[str] | None = None,
    infer_fks: bool = True,
) -> tuple[MetaData, list[tuple[str, str, str, str]], list[tuple[str, str]]]:
    engine = create_engine("sqlite:///:memory:")
    meta = MetaData()
    included_tables = set(table_filter) if table_filter else set(tables.keys())
    pk_map = choose_primary_columns(tables)
    inferred, unresolved = infer_relationships(tables, included_tables) if infer_fks else ([], [])
    rel_by_source = {(st, sc): (tt, tc) for st, sc, tt, tc in inferred}

    for tname in sorted(included_tables):
        table = tables.get(tname)
        if not table:
            print(f"    [WARN] table '{tname}' listed for ERD but no CSV file was found")
            continue

        cols: list[Column] = []
        for col in table.columns:
            args = [_sa_type(col.sql_type)]
            kwargs = {}
            target = rel_by_source.get((tname, col.name))
            if target:
                target_table, target_col = target
                target_safe_table = tables[target_table].safe_name
                target_safe_col = next(c.safe_name for c in tables[target_table].columns if c.name == target_col)
                args.append(ForeignKey(f"{target_safe_table}.{target_safe_col}"))
            if pk_map.get(tname) == col.name:
                kwargs["primary_key"] = True
            cols.append(Column(col.safe_name, *args, **kwargs))

        try:
            Table(table.safe_name, meta, *cols)
        except Exception as exc:
            print(f"    [WARN] {tname}: skipped; could not build table: {exc}")

    try:
        meta.create_all(engine)
    except Exception:
        pass
    return meta, inferred, unresolved


def isolated_tables(
    table_names: set[str],
    relationships: list[tuple[str, str, str, str]],
) -> list[str]:
    connected: set[str] = set()
    for source_table, _source_col, target_table, _target_col in relationships:
        connected.update({source_table, target_table})
    return sorted(table_names - connected)


def build_table_alias_map(tables: dict[str, CsvTable]) -> dict[str, str]:
    alias_map: dict[str, str] = {}

    def add(alias: str, table_name: str) -> None:
        norm = _normalize(alias)
        if norm and norm not in alias_map:
            alias_map[norm] = table_name

    for table_name in tables:
        lower = table_name.lower()
        no_ext = lower.removesuffix(".csv")
        add(table_name, table_name)
        add(lower, table_name)
        add(no_ext, table_name)
        add(no_ext.replace("_", ""), table_name)
        add(_singular(no_ext), table_name)
        add(_singular(no_ext.replace("_", "")), table_name)

        if no_ext == "inventory":
            add("inv", table_name)
            add("inv_df", table_name)
        elif no_ext == "products":
            add("product", table_name)
            add("prod", table_name)
            add("prod_df", table_name)
        elif no_ext == "customers":
            add("customer", table_name)
            add("cust", table_name)
            add("cust_df", table_name)
        elif no_ext == "cust_hist":
            add("custhist", table_name)
            add("custhist_df", table_name)
            add("customer_history", table_name)
        elif no_ext == "orderlines":
            add("orderline", table_name)
            add("orderlines_df", table_name)
        elif no_ext == "orders":
            add("order", table_name)
            add("orders_df", table_name)

    return alias_map


def resolve_domain_tables(
    raw_table_names: list[str],
    tables: dict[str, CsvTable],
    alias_map: dict[str, str],
) -> tuple[list[str], list[str], list[tuple[str, str]]]:
    resolved: list[str] = []
    seen: set[str] = set()
    missing: list[str] = []
    aliases: list[tuple[str, str]] = []

    for raw_name in raw_table_names:
        table_name = alias_map.get(_normalize(raw_name))
        if table_name is None and raw_name in tables:
            table_name = raw_name
        if table_name is None:
            missing.append(raw_name)
            continue
        if raw_name != table_name:
            aliases.append((raw_name, table_name))
        if table_name not in seen:
            resolved.append(table_name)
            seen.add(table_name)

    return resolved, missing, aliases


def read_best_run_tag(dataset_dir: Path) -> str | None:
    xlsx = dataset_dir / "ccm_output" / "tune_params_results.xlsx"
    if not xlsx.exists():
        return None
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        ws = wb.active
        row = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
        return str(row[3]) if row[3] is not None else None
    except Exception as exc:
        print(f"  [WARN] Could not read {xlsx}: {exc}")
        return None


def render_erd(meta: MetaData, out_path: Path, title: str, fmt: str) -> bool:
    if not list(meta.tables):
        print(f"    [SKIP] No tables; skipping {out_path.name}")
        return False
    target = str(out_path.with_suffix(f".{fmt}"))
    try:
        render_er(meta, target, title=title)
        if not Path(target).exists():
            print(f"    [ERR] {out_path.name}: render_er returned, but no file was written")
            return False
        print(f"    OK {out_path.name} ({len(list(meta.tables))} tables)")
        return True
    except Exception as exc:
        import traceback

        print(f"    [ERR] {out_path.name}: {type(exc).__name__}: {exc}")
        traceback.print_exc()
        return False


def load_domains(dataset_dir: Path) -> tuple[str | None, list[dict]]:
    run_tag = read_best_run_tag(dataset_dir)
    if not run_tag:
        return None, []
    domains_path = dataset_dir / "ccm_output" / run_tag / "step5_domains.json"
    if not domains_path.exists():
        print(f"  [WARN] step5_domains.json not found at {domains_path}; domain ERDs skipped")
        return run_tag, []
    try:
        return run_tag, json.loads(domains_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception as exc:
        print(f"  [WARN] Cannot read step5_domains.json: {exc}; domain ERDs skipped")
        return run_tag, []


def process_dataset(
    dataset_dir: Path,
    erd_root: Path,
    fmt: str,
    infer_fks: bool = True,
    max_distinct: int = DEFAULT_MAX_DISTINCT,
) -> tuple[int, int]:
    name = dataset_dir.name
    print(f"\n  {'-' * 50}")
    print(f"  Dataset: {name}")

    csv_dir = dataset_dir / "csv"
    if not csv_dir.exists():
        print(f"  [SKIP] csv folder not found: {csv_dir}")
        return (0, 0)

    try:
        tables = read_csv_tables(csv_dir, max_distinct=max_distinct)
    except Exception as exc:
        print(f"  [SKIP] Cannot read CSV files: {exc}")
        return (0, 0)
    if not tables:
        print("  [SKIP] no CSV tables found")
        return (0, 0)
    alias_map = build_table_alias_map(tables)

    run_tag, domains = load_domains(dataset_dir)
    print(f"  Source   : {csv_dir}")
    print(f"  Tables   : {len(tables)} CSV table(s)")
    print(f"  Best run : {run_tag or 'not found'}")
    print(f"  Domains  : {len(domains)}")

    dataset_out = erd_root / name
    domains_out = dataset_out / "domains"
    dataset_out.mkdir(parents=True, exist_ok=True)
    domains_out.mkdir(parents=True, exist_ok=True)

    n_ok, n_fail = 0, 0

    print("  Rendering full-schema ERD from CSV files ...")
    full_meta, inferred, unresolved = build_metadata_from_csv(tables, infer_fks=infer_fks)
    if infer_fks:
        print(f"    CSV relationships inferred: {len(inferred)} (unresolved key-like columns: {len(unresolved)})")
    if render_erd(full_meta, dataset_out / "full_schema", f"{name} - Full Schema ({len(full_meta.tables)} tables)", fmt):
        n_ok += 1
    else:
        n_fail += 1

    print(f"  Rendering {len(domains)} domain ERD(s) from related CSV tables ...")
    domain_summaries: list[tuple[str, str, int, int, list[str]]] = []
    domain_aliases: list[tuple[str, str, str]] = []
    for domain in domains:
        did = domain.get("domain_id", "?")
        dname = domain.get("domain_name", f"Domain_{did}")
        raw_domain_tables = list(domain.get("tables", []))
        domain_tables, missing, aliases = resolve_domain_tables(raw_domain_tables, tables, alias_map)
        domain_aliases.extend((str(did), raw, resolved) for raw, resolved in aliases)
        if missing:
            print(f"    [WARN] D{did}: {len(missing)} domain table(s) missing CSV: {missing[:5]}")
        if len(domain_tables) != len(raw_domain_tables):
            print(f"    D{did}: {len(raw_domain_tables)} domain name(s) resolved to {len(domain_tables)} CSV table(s)")
        if not domain_tables:
            continue

        domain_meta, domain_inferred, _ = build_metadata_from_csv(
            tables,
            table_filter=domain_tables,
            infer_fks=infer_fks,
        )
        if infer_fks:
            domain_summaries.append(
                (
                    str(did),
                    str(dname),
                    len(domain_tables),
                    len(domain_inferred),
                    isolated_tables(set(domain_tables), domain_inferred),
                )
            )
        fname = f"D{did}_{_slug(dname)}"
        title = f"{name} - D{did}: {dname} ({len(domain_tables)} tables)"
        if render_erd(domain_meta, domains_out / fname, title, fmt):
            n_ok += 1
        else:
            n_fail += 1

    if infer_fks:
        log_path = dataset_out / "inferred_relationships.txt"
        full_isolated = isolated_tables(set(tables), inferred)
        lines = [
            f"CSV relationship inference for {name}",
            "(heuristic: source column distinct values are contained in a unique target column)",
            "=" * 70,
            "",
            f"Inferred ({len(inferred)}):",
        ]
        for source_table, source_col, target_table, target_col in sorted(set(inferred)):
            lines.append(f"  {source_table}.{source_col} -> {target_table}.{target_col}")
        lines += ["", f"Unresolved key-like source columns ({len(unresolved)}):"]
        for source_table, source_col in sorted(set(unresolved)):
            lines.append(f"  {source_table}.{source_col}")
        lines += ["", f"Isolated full-schema tables ({len(full_isolated)}):"]
        for table_name in full_isolated:
            lines.append(f"  {table_name}")
        lines += ["", "Domain relationship summaries:"]
        for did, dname, table_count, rel_count, domain_isolated in domain_summaries:
            lines.append(f"  D{did} {dname}: {table_count} table(s), {rel_count} inferred relationship(s)")
            if domain_isolated:
                lines.append(f"    isolated: {', '.join(domain_isolated)}")
        if domain_aliases:
            lines += ["", "Domain table aliases resolved to CSV tables:"]
            for did, raw_name, resolved_name in domain_aliases:
                lines.append(f"  D{did}: {raw_name} -> {resolved_name}")
        log_path.write_text("\n".join(lines), encoding="utf-8")

    if n_fail:
        print(f"  WARN {name}: {n_ok} rendered, {n_fail} failed")
    return n_ok, n_fail


def discover_datasets(root: Path) -> list[str]:
    found = []
    for child in sorted(root.iterdir()):
        if child.is_dir() and (child / "csv").is_dir():
            found.append(child.name)
    return found


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate ERDs directly from dataset CSV files.")
    parser.add_argument("--root", type=Path, default=DATALAKES_ROOT, help="Dataset root (default: Datalakes/)")
    parser.add_argument("--output_dir", type=Path, default=None, help="ERD output root (default: project ERDs/)")
    parser.add_argument("--datasets", nargs="+", default=None, help="Specific datasets to process")
    parser.add_argument("--format", choices=["png", "pdf", "svg"], default="png", help="Output format")
    parser.add_argument("--no-infer-fks", action="store_true", help="Do not infer FK relationship lines")
    parser.add_argument(
        "--max-distinct",
        type=int,
        default=DEFAULT_MAX_DISTINCT,
        help=f"Maximum distinct values stored per column for relationship inference (default: {DEFAULT_MAX_DISTINCT})",
    )
    args = parser.parse_args()

    root = args.root.resolve()
    erd_root = (args.output_dir or PROJECT_ROOT / "ERDs").resolve()
    datasets = args.datasets or discover_datasets(root)
    infer_fks = not args.no_infer_fks

    print("\nDomainMiner - CSV ERD Generator")
    print(f"Root       : {root}")
    print(f"ERD output : {erd_root}")
    print(f"Format     : {args.format}")
    print(f"FK lines   : {'inferred from CSV values' if infer_fks else 'disabled'}")
    print(f"Datasets   : {len(datasets)} candidate(s)")

    check_graphviz()
    erd_root.mkdir(parents=True, exist_ok=True)

    total_ok, total_fail, skipped = 0, 0, 0
    for dataset_name in datasets:
        dataset_dir = root / dataset_name
        if not dataset_dir.is_dir():
            print(f"\n  [SKIP] {dataset_name}: folder not found")
            skipped += 1
            continue
        try:
            n_ok, n_fail = process_dataset(
                dataset_dir,
                erd_root,
                args.format,
                infer_fks=infer_fks,
                max_distinct=args.max_distinct,
            )
            total_ok += n_ok
            total_fail += n_fail
            if n_ok == 0 and n_fail == 0:
                skipped += 1
        except Exception as exc:
            import traceback

            print(f"\n  [ERR] {dataset_name}: {exc}")
            traceback.print_exc()
            skipped += 1

    print(f"\n{'-' * 52}")
    print(f"Done. {total_ok} ERD(s) rendered, {total_fail} failed, {skipped} dataset(s) skipped.")
    print(f"ERDs written to: {erd_root}")


if __name__ == "__main__":
    main()
